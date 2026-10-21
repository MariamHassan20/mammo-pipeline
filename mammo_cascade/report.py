"""
Organised outputs for a run:

  <out>/segmented_tumors/   <name>_patch.png (tumour crop) and <name>_mask.png (tumour mask)
  <out>/overlays/           <name>_overlay.png
  <out>/json/               <name>.json
  <out>/diagnosis.xlsx      one row per image
  <out>/metrics.xlsx        only with ground truth: accuracy, AUC, F1, ... and segmentation Dice/IoU
  <out>/examples/           only with ground-truth masks: 4 cases, our segmentation vs the true tumour
"""
import os
import json
import datetime
import numpy as np
import cv2
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from .models import cap_size

CLASSES = ("normal", "benign", "malignant")


# ------------------------------------------------------------------ ground truth
def load_labels(path):
    """CSV with columns image, label (normal/benign/malignant or 0=benign/1=malignant), optional mask."""
    import csv
    norm = {"0": "benign", "1": "malignant", "cancer": "malignant"}
    out, by_name = {}, {}
    with open(path, newline="", encoding="utf-8-sig") as f:      # utf-8-sig: CSVs saved by Excel
        reader = csv.DictReader(f)
        if not reader.fieldnames or not {"image", "label"} <= set(reader.fieldnames):
            raise ValueError(f"labels file needs the columns 'image' and 'label'; found {reader.fieldnames}")
        for r in reader:
            if not (r.get("image") or "").strip():
                continue                                         # blank line
            lab = str(r.get("label") or "").strip().lower()
            lab = norm.get(lab, lab)
            if lab not in CLASSES:
                raise ValueError(f"unknown label '{r.get('label')}' for {r.get('image')}")
            img = r["image"].strip()
            rec = {"label": lab, "mask": (r.get("mask") or "").strip() or None}
            out[os.path.abspath(img)] = rec
            by_name.setdefault(os.path.basename(img), []).append(rec)
    # a bare file name is only used when it is unique in the CSV (same names in different folders are ambiguous)
    for name, recs in by_name.items():
        if len(recs) == 1:
            out.setdefault(name, recs[0])
    return out


def lookup(labels, image_path):
    return labels.get(os.path.abspath(image_path)) or labels.get(os.path.basename(image_path))


def gt_mask(path, shape):
    """Ground-truth mask resized to the working image (same 4000-px cap as the pipeline)."""
    m = cv2.imread(path, cv2.IMREAD_GRAYSCALE | cv2.IMREAD_ANYDEPTH)
    if m is None:
        return None
    m, _ = cap_size(m)
    if m.shape != shape:
        m = cv2.resize(m, (shape[1], shape[0]), interpolation=cv2.INTER_NEAREST)
    m = (m > (0 if m.max() <= 1 else m.max() / 2)).astype(np.uint8)   # masks stored as 0/1, 0/255 or 16-bit
    return m if m.any() else None


def seg_scores(pred, gt):
    p = pred.astype(bool) if pred is not None else np.zeros(gt.shape, bool)
    g = gt.astype(bool)
    inter, ps, gs = (p & g).sum(), p.sum(), g.sum()
    return float(2 * inter / (ps + gs)), float(inter / max((p | g).sum(), 1))


# ------------------------------------------------------------------ per-image files
def save_image_outputs(result, art, out_dir, name):
    """Write overlay, tumour patch and mask; return their relative paths."""
    from .visualize import draw_overlay
    paths = {"overlay": os.path.join("overlays", f"{name}_overlay.png")}
    cv2.imwrite(os.path.join(out_dir, paths["overlay"]), draw_overlay(result, art))
    g = art["gray"]
    if result["stage_a"]["passed"]:              # normal images get no tumour patch or mask
        if art.get("crop_box") is not None:
            x0, y0, x1, y1 = art["crop_box"]
            paths["patch"] = os.path.join("segmented_tumors", f"{name}_patch.png")
            cv2.imwrite(os.path.join(out_dir, paths["patch"]), g[y0:y1, x0:x1])
        if art.get("mask") is not None and art["mask"].any():
            paths["mask"] = os.path.join("segmented_tumors", f"{name}_mask.png")
            cv2.imwrite(os.path.join(out_dir, paths["mask"]), art["mask"] * 255)
    with open(os.path.join(out_dir, "json", f"{name}.json"), "w") as f:
        json.dump(result, f, indent=2)
    return paths


def make_dirs(out_dir):
    for d in ("segmented_tumors", "overlays", "json"):
        os.makedirs(os.path.join(out_dir, d), exist_ok=True)


# ------------------------------------------------------------------ excel helpers
_HEAD = Font(bold=True, color="FFFFFF")
_FILL = PatternFill("solid", fgColor="305496")


def _sheet(wb, title, header, rows, first=False):
    ws = wb.active if first else wb.create_sheet()
    ws.title = title
    ws.append(header)
    for c in ws[1]:
        c.font, c.fill = _HEAD, _FILL
    for r in rows:
        ws.append(r)
    for i, h in enumerate(header, 1):
        width = max([len(str(h))] + [len(str(r[i - 1])) for r in rows if i - 1 < len(r) and r[i - 1] is not None])
        ws.column_dimensions[get_column_letter(i)].width = min(60, width + 2)
    ws.freeze_panes = "A2"
    return ws


DIAG_COLS = ["image", "decision", "true_label", "correct", "p_normal", "p_benign", "p_malignant",
             "p_abnormal", "stage_a_passed", "p_malignant_patch", "p_malignant_whole", "box_source", "box_score",
             "box_xyxy", "lesion_area_px", "equiv_diameter_px", "area_mm2", "equiv_diameter_mm", "dice", "iou",
             "patch_file", "mask_file", "overlay_file", "warnings", "error", "seconds"]


def diagnosis_row(image, result, paths, truth=None, dice=None, iou=None, error=""):
    if result is None:
        return {"image": image, "decision": "ERROR", "true_label": truth, "error": error}
    a, sb = result["stage_a"], result.get("stage_b", {})
    loc, seg = result.get("localisation", {}), result.get("segmentation", {})
    stopped = not a["passed"]
    ev = result.get("evaluation_only", {})
    # images stopped at Stage A show no later-step values; those are kept only for the metrics
    hidden = {"p_malignant_patch": ev.get("stage_b", {}).get("p_malignant_patch"),
              "p_malignant_whole": ev.get("stage_b", {}).get("p_malignant_whole"),
              "box_source": ev.get("localisation", {}).get("source"),
              "dice": None if dice is None else round(dice, 4), "iou": None if iou is None else round(iou, 4)}
    if stopped:
        dice = iou = None
    return {
        "_eval": hidden if stopped else {},
        "image": image, "decision": result["decision"], "true_label": truth,
        "correct": (result["decision"] == truth) if truth else None,
        "p_normal": a["probs"].get("normal"), "p_benign": a["probs"].get("benign"),
        "p_malignant": a["probs"].get("malignant"), "p_abnormal": a["p_abnormal"], "stage_a_passed": a["passed"],
        "p_malignant_patch": sb.get("p_malignant_patch"), "p_malignant_whole": sb.get("p_malignant_whole"),
        "box_source": loc.get("source"), "box_score": loc.get("score"),
        "box_xyxy": str(loc["box_xyxy"]) if loc.get("box_xyxy") else None,
        "lesion_area_px": seg.get("area_px"), "equiv_diameter_px": seg.get("equiv_diameter_px"),
        "area_mm2": seg.get("area_mm2"), "equiv_diameter_mm": seg.get("equiv_diameter_mm"),
        "dice": None if dice is None else round(dice, 4), "iou": None if iou is None else round(iou, 4),
        "patch_file": paths.get("patch"), "mask_file": paths.get("mask"), "overlay_file": paths.get("overlay"),
        "warnings": "; ".join(result.get("warnings", [])), "error": error, "seconds": result.get("seconds"),
    }


def write_diagnosis(rows, path):
    wb = Workbook()
    _sheet(wb, "diagnosis", DIAG_COLS, [[r.get(c) for c in DIAG_COLS] for r in rows], first=True)
    wb.save(path)


# ------------------------------------------------------------------ metrics
def _auc(y, s):
    from sklearn.metrics import roc_auc_score
    return float(roc_auc_score(y, s)) if len(set(y)) == 2 else None


def _auc_ci(y, s, n=1000, seed=0):
    rng = np.random.default_rng(seed)
    y, s = np.asarray(y), np.asarray(s)
    vals = []
    for _ in range(n):
        i = rng.integers(0, len(y), len(y))
        if len(set(y[i])) == 2:
            vals.append(_auc(y[i], s[i]))
    return (round(float(np.percentile(vals, 2.5)), 4), round(float(np.percentile(vals, 97.5)), 4)) if vals else (None, None)


def _binary(y, p):
    y, p = np.asarray(y, bool), np.asarray(p, bool)
    tp, tn = int((y & p).sum()), int((~y & ~p).sum())
    fp, fn = int((~y & p).sum()), int((y & ~p).sum())
    sens = tp / (tp + fn) if tp + fn else None
    spec = tn / (tn + fp) if tn + fp else None
    prec = tp / (tp + fp) if tp + fp else None
    f1 = 2 * prec * sens / (prec + sens) if prec and sens else 0.0
    return {"accuracy": (tp + tn) / len(y), "balanced_accuracy": None if sens is None or spec is None else (sens + spec) / 2,
            "f1": f1, "precision": prec, "sensitivity_recall": sens, "specificity": spec,
            "TP": tp, "FP": fp, "TN": tn, "FN": fn}


def _val(r, k):
    """Value for the metrics: the visible one, or the evaluation-only one for images stopped at Stage A."""
    v = r.get(k)
    return v if v is not None else r.get("_eval", {}).get(k)


def write_metrics(rows, path, settings):
    ok = [r for r in rows if r.get("true_label") and r["decision"] != "ERROR"]
    wb = Workbook()
    if not ok:
        _sheet(wb, "classification", ["metric", "value"],
               [["images with ground truth", 0], ["note", "no input image matched a row of the labels file"]],
               first=True)
        wb.save(path)
        return
    R = lambda v: None if v is None else round(float(v), 4)
    cls_rows = [["images with ground truth", len(ok), ""]]
    for c in CLASSES:
        cls_rows.append([f"  true {c}", sum(r["true_label"] == c for r in ok), ""])

    # final diagnosis vs truth, malignant = positive
    y = [r["true_label"] == "malignant" for r in ok]
    pred = [r["decision"] == "malignant" for r in ok]
    b = _binary(y, pred)
    cls_rows.append(["FINAL DIAGNOSIS (malignant = positive; 'normal' and 'benign' count as negative)", "", ""])
    for k in ("accuracy", "balanced_accuracy", "f1", "precision", "sensitivity_recall", "specificity"):
        cls_rows.append([f"  {k}", R(b[k]), ""])
    for k in ("TP", "FP", "TN", "FN"):
        cls_rows.append([f"  {k}", b[k], ""])
    # 3-class accuracy of the final decision
    cls_rows.append(["  exact-match accuracy (normal/benign/malignant)", R(np.mean([r["correct"] for r in ok])), ""])

    # AUCs on continuous scores (computed for every image, even if Stage A said normal)
    cls_rows.append(["AUC (benign vs malignant, images with both labels)", "value", "95% CI (bootstrap)"])
    bm = [r for r in ok if r["true_label"] in ("benign", "malignant")]
    yb = [r["true_label"] == "malignant" for r in bm]
    for key, label in (("p_malignant_patch", "Stage B on tumour patch"), ("p_malignant_whole", "Stage B on whole image"),
                       ("p_malignant", "Stage A P(malignant)")):
        sc = [_val(r, key) for r in bm]
        if bm and all(v is not None for v in sc):
            lo, hi = _auc_ci(yb, sc)
            cls_rows.append([f"  {label}", R(_auc(yb, sc)), f"{lo} - {hi}"])
    # Stage A gate
    ya = [r["true_label"] != "normal" for r in ok]
    cls_rows.append(["STAGE A GATE (abnormal = positive)", "", ""])
    if len(set(ya)) == 2:
        cls_rows.append(["  AUC P(abnormal)", R(_auc(ya, [r["p_abnormal"] for r in ok])), ""])
    for c in ("benign", "malignant", "normal"):
        sub = [r for r in ok if r["true_label"] == c]
        if sub:
            cls_rows.append([f"  passed to segmentation: true {c}", R(np.mean([bool(r["stage_a_passed"]) for r in sub])), ""])
    _sheet(wb, "classification", ["metric", "value", "note"], cls_rows, first=True)

    # confusion table
    conf = [[f"true {t}"] + [sum(r["true_label"] == t and r["decision"] == p for r in ok) for p in CLASSES]
            for t in CLASSES if any(r["true_label"] == t for r in ok)]
    _sheet(wb, "confusion", ["", "pred normal", "pred benign", "pred malignant"], conf)

    # segmentation
    # segmentation is scored on every image with a true mask, including those stopped at Stage A
    seg = [r for r in ok if _val(r, "dice") is not None]
    if seg:
        d, j = np.array([_val(r, "dice") for r in seg]), np.array([_val(r, "iou") for r in seg])
        srows = [["images with a ground-truth mask (all, including those stopped at Stage A)", len(seg)],
                 ["Dice mean", R(d.mean())], ["Dice median", R(np.median(d))], ["Dice std", R(d.std())],
                 ["IoU mean", R(j.mean())], ["IoU median", R(np.median(j))], ["fraction Dice >= 0.5", R((d >= 0.5).mean())]]
        for src in sorted({_val(r, "box_source") or "none" for r in seg}):
            ds = np.array([_val(r, "dice") for r in seg if (_val(r, "box_source") or "none") == src])
            srows.append([f"Dice mean, box from {src} (n={len(ds)})", R(ds.mean())])
        _sheet(wb, "segmentation", ["metric", "value"], srows)

    _sheet(wb, "settings", ["setting", "value"], [[k, str(v)] for k, v in settings.items()] +
           [["created", datetime.datetime.now().isoformat(timespec="seconds")]])
    wb.save(path)


# ------------------------------------------------------------------ examples
_FONT, _FS = cv2.FONT_HERSHEY_SIMPLEX, 0.6


def _tile(img, h=400):
    s = h / img.shape[0]
    return cv2.resize(img, (max(1, int(img.shape[1] * s)), h), interpolation=cv2.INTER_AREA)


def _label(img, text):
    """Add a caption above a tile; widen the tile with white margins if the caption is longer."""
    tw = cv2.getTextSize(text, _FONT, _FS, 1)[0][0] + 12
    if img.shape[1] < tw:
        extra = tw - img.shape[1]
        img = cv2.copyMakeBorder(img, 0, 0, extra // 2, extra - extra // 2, cv2.BORDER_CONSTANT, value=(255, 255, 255))
    out = cv2.copyMakeBorder(img, 34, 0, 0, 0, cv2.BORDER_CONSTANT, value=(255, 255, 255))
    cv2.putText(out, text, (6, 24), _FONT, _FS, (0, 0, 0), 1, cv2.LINE_AA)
    return out


def _header(lines, width):
    width = max([width] + [cv2.getTextSize(t, _FONT, _FS, 1)[0][0] + 12 for t in lines])
    head = np.full((30 * len(lines) + 8, width, 3), 255, np.uint8)
    for i, t in enumerate(lines):
        cv2.putText(head, t, (6, 24 + 30 * i), _FONT, _FS, (0, 0, 0), 1, cv2.LINE_AA)
    return head


def save_example(gray, pred_mask, gt, crop_box, row, path, show_numbers=True):
    """Panel: full breast | region with true (green) and predicted (red) outline | our tumour patch | true tumour patch.
    show_numbers=False prints only the colour legend (no labels, scores or Dice)."""
    ys, xs = np.where(gt > 0)
    gx0, gy0, gx1, gy1 = xs.min(), ys.min(), xs.max(), ys.max()
    pad = int(0.10 * max(gx1 - gx0, gy1 - gy0) + 5)
    H, W = gray.shape
    tb = [max(0, gx0 - pad), max(0, gy0 - pad), min(W, gx1 + pad), min(H, gy1 + pad)]   # true-tumour patch
    boxes = [tb] + ([crop_box] if crop_box else [])
    vx0, vy0 = min(b[0] for b in boxes), min(b[1] for b in boxes)
    vx1, vy1 = max(b[2] for b in boxes), max(b[3] for b in boxes)
    m = int(0.15 * max(vx1 - vx0, vy1 - vy0))
    vx0, vy0, vx1, vy1 = max(0, vx0 - m), max(0, vy0 - m), min(W, vx1 + m), min(H, vy1 + m)
    full = cv2.cvtColor(gray, cv2.COLOR_GRAY2BGR)
    reg = full[vy0:vy1, vx0:vx1].copy()
    th = max(2, int(0.006 * max(reg.shape[:2])))
    th_full = max(3, int(0.004 * max(H, W)))
    for mk, col in ((gt, (0, 200, 0)), (pred_mask, (0, 0, 255))):
        if mk is not None and mk.any():
            c, _ = cv2.findContours(mk.astype(np.uint8), cv2.RETR_EXTERNAL, cv2.CHAIN_APPROX_SIMPLE)
            cv2.drawContours(full, c, -1, col, th_full)
            c, _ = cv2.findContours(mk[vy0:vy1, vx0:vx1].astype(np.uint8), cv2.RETR_EXTERNAL, cv2.CHAIN_APPROX_SIMPLE)
            cv2.drawContours(reg, c, -1, col, th)
    cv2.rectangle(full, (vx0, vy0), (vx1, vy1), (0, 200, 255), th_full)       # yellow = zoomed region
    tiles = [_label(_tile(full, 600), "full breast"), _label(_tile(reg, 600), "tumour region")]
    if crop_box:
        x0, y0, x1, y1 = crop_box
        tiles.append(_label(_tile(cv2.cvtColor(gray[y0:y1, x0:x1], cv2.COLOR_GRAY2BGR), 600), "our tumour patch"))
    tiles.append(_label(_tile(cv2.cvtColor(gray[tb[1]:tb[3], tb[0]:tb[2]], cv2.COLOR_GRAY2BGR), 600), "true tumour patch"))
    sep = np.full((tiles[0].shape[0], 10, 3), 255, np.uint8)
    body = tiles[0]
    for t in tiles[1:]:
        body = np.hstack([body, sep, t])
    legend = "GREEN = true tumour (ground truth), RED = our segmentation, YELLOW box = zoomed region"
    lines = ([f"true: {row['true_label']}    predicted: {row['decision']}    "
              f"P(malignant, patch) = {row.get('p_malignant_patch')}",
              f"segmentation Dice = {row.get('dice')}    box from: {row.get('box_source')}", legend]
             if show_numbers else [legend])
    head = _header(lines, body.shape[1])
    if head.shape[1] > body.shape[1]:
        body = cv2.copyMakeBorder(body, 0, 0, 0, head.shape[1] - body.shape[1], cv2.BORDER_CONSTANT, value=(255, 255, 255))
    cv2.imwrite(path, np.vstack([head, body]))
